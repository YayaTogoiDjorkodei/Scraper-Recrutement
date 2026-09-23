import json
from html import escape

from openpyxl import load_workbook
from corporate_scraper.extraction import extract_structured_requirements, requirement_summary
from corporate_scraper.exports import export_study
from corporate_scraper.models import JobObservation, RunSpec
from corporate_scraper.paths import results_directory, set_results_directory
from corporate_scraper.references import TECHNOLOGY_REFERENCES
from corporate_scraper.reprocess import reprocess_study
from corporate_scraper.sources import LinkedInAdapter
from corporate_scraper.store import StudyStore
from corporate_scraper.text import description_text


def test_escaped_description_reaches_excel_and_survives_reprocessing(tmp_path):
    raw = '<p>Bac+5 en informatique.</p><p>Expérience : 1 à 2 ans en développement.</p><p>C# et ASP.NET Core, SQL requis. CDI.</p>'
    page = '<script type="application/ld+json">' + json.dumps({"@type": "JobPosting", "description": escape(escape(raw))}) + '</script>'
    text = LinkedInAdapter().parse_description(page)
    assert '<p>' not in text and '\n' in text
    store = StudyStore(tmp_path / 'studies.db')
    run_id = store.create_run(RunSpec('Pipeline', ('developer',), ('Rabat',), ('linkedin',)))
    store.commit_discovery(run_id, 'page1', [JobObservation('linkedin', '1', 'https://example.test/1', 'Developer')])
    # Simulate the old broken record, including its unprocessed source text.
    with store._connection() as connection:
        connection.execute("UPDATE offers SET description=?, description_status='available'", (raw,))
    assert reprocess_study(store, run_id) == 1
    record = store.offers(run_id)[0]
    assert record['raw_description'] == raw
    assert record['education'] == 'Bac+5'
    assert '1 à 2 ans' in record['experience']
    assert record['contract'] == 'CDI'
    assert {'C#', 'SQL'} <= set(json.loads(record['skills_json']))
    output = export_study(store, run_id, tmp_path / 'output.xlsx', True)
    book = load_workbook(output)
    assert 'C#' in book['IT Jobs Data']['D2'].value
    assert book['IT Jobs Data']['E2'].value == 'Bac+5'
    assert '1 à 2 ans' in book['IT Jobs Data']['F2'].value
    assert '<p>' not in book['Descriptions']['E2'].value
    assert book['IT Jobs Data'].row_dimensions[2].height <= 70
    assert reprocess_study(store, run_id) == 1
    assert store.offers(run_id)[0]['raw_description'] == raw


def test_english_requirements_and_negative_education_cases():
    text = "Bachelor's degree in computer science. At least 3-5 years of experience. Python and Docker required. Kubernetes preferred. Scrum Master works with master data."
    fields = requirement_summary(extract_structured_requirements(text, TECHNOLOGY_REFERENCES))
    assert "Bachelor's degree" in fields['education']
    assert 'master' not in fields['education'].lower()
    assert '3-5 years' in fields['experience']
    assert {'Python', 'Docker', 'Kubernetes'} <= set(fields['skills'])
    negative = requirement_summary(extract_structured_requirements("Ingénieur réseau. Notre société existe depuis 30 ans.", TECHNOLOGY_REFERENCES))
    assert negative['education'] == negative['experience'] == ''


def test_html_inline_tags_do_not_split_skills():
    assert description_text('<p><strong>ASP</strong>.NET &amp; C#</p><p>Bac+5</p>') == 'ASP.NET & C#\nBac+5'


def test_common_prose_does_not_become_skills():
    fields = requirement_summary(extract_structured_requirements('Un chef accompagne ses équipes. Solid experience in development and rest days.', TECHNOLOGY_REFERENCES))
    assert not fields['skills']
    fields = requirement_summary(extract_structured_requirements('Amazon SES, SOLID and REST APIs. Vue and Vue.js.', TECHNOLOGY_REFERENCES))
    assert {'Amazon SES', 'SOLID', 'REST', 'Vue.js'} <= set(fields['skills'])
    assert 'Vue' not in fields['skills']
    fields = requirement_summary(extract_structured_requirements('Vue.js and Node.js development.', TECHNOLOGY_REFERENCES))
    assert 'JavaScript' not in fields['skills']


def test_company_history_and_technical_acronyms_are_not_job_requirements():
    text = "With over 40 years of experience we provide innovation at every stage of the value chain. Verify AUTOSAR MCAL, CDD and BSW software. Première expérience (stage ou projet) en Python."
    fields = requirement_summary(extract_structured_requirements(text, TECHNOLOGY_REFERENCES))
    assert fields['experience'] == ''
    assert fields['contract'] == ''
    fields = requirement_summary(extract_structured_requirements('Stage rémunéré de fin d’études. Type de contrat : CDI à temps plein.', TECHNOLOGY_REFERENCES))
    assert 'Stage' in fields['contract'] and 'CDI' in fields['contract']
    assert 'temps plein' not in fields['contract']


def test_results_default_and_persisted_override(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    assert results_directory() == tmp_path / 'results'
    chosen = set_results_directory(str(tmp_path / 'my-output'))
    assert chosen == results_directory() == tmp_path / 'my-output'


def test_long_description_is_preserved_and_formula_values_are_text(tmp_path):
    store = StudyStore(tmp_path / 'studies.db')
    run_id = store.create_run(RunSpec('Long', ('python',), ('Rabat',), ('linkedin',)))
    text = 'Python experience and documentation. ' * 1200
    store.commit_discovery(run_id, 'search', [JobObservation('linkedin', '1', 'https://example.test/1', '=Danger')])
    store.commit_detail(run_id, '1', text, 'available', 'extracted', ())
    book = load_workbook(export_study(store, run_id, tmp_path / 'long.xlsx', True))
    recovered = ''.join(row[4] for row in book['Descriptions'].iter_rows(min_row=2, values_only=True))
    assert recovered == description_text(text)
    assert book['IT Jobs Data']['B2'].data_type == 's'


def test_settings_and_automatic_export_use_the_same_saved_folder(qtbot, tmp_path, monkeypatch):
    from corporate_scraper.ui import SettingsPage, CollectionPage
    monkeypatch.chdir(tmp_path)
    settings = SettingsPage('studies.sqlite3')
    qtbot.addWidget(settings)
    settings.results_path.setText(str(tmp_path / 'chosen'))
    settings._save()
    reopened = SettingsPage('studies.sqlite3')
    qtbot.addWidget(reopened)
    assert reopened.results_path.text() == str(tmp_path / 'chosen')
    store = StudyStore(tmp_path / 'studies.sqlite3')
    run_id = store.create_run(RunSpec('UI export', ('python',), ('Rabat',), ('linkedin',)))
    store.commit_discovery(run_id, 'page1', [JobObservation('linkedin', '1', 'https://example.test/1', 'Developer')])
    page = CollectionPage()
    qtbot.addWidget(page)
    page._store, page._run_id = store, run_id
    page._start_automatic_export()
    qtbot.waitUntil(lambda: page.open_excel_button.isEnabled(), timeout=10000)
    assert (tmp_path / 'chosen' / f'etude-{run_id[:8]}.xlsx').exists()
    page._export_thread.wait(1000)
