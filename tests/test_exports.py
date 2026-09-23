from openpyxl import load_workbook

from corporate_scraper.exports import export_study
from corporate_scraper.models import FieldEvidence, JobObservation, RunSpec
from corporate_scraper.store import StudyStore


def test_default_workbook_has_safe_values_links_and_evidence(tmp_path):
    store = StudyStore(tmp_path / "studies.db")
    run_id = store.create_run(RunSpec(name="Export", queries=("python",), cities=("Rabat",), sources=("linkedin",)))
    offer = JobObservation(source="linkedin", source_key="safe", canonical_url="https://example.test/job", title="=Safe title",
                           company="Acme", location="Rabat", skills=("Python",), description_status="pending")
    store.commit_discovery(run_id, "search:1", [offer])
    store.commit_detail(run_id, "safe", "Description complète", "available", "verified",
                        (FieldEvidence("skill_required", "Python", "Python requis", "exact", 100),))
    store.commit_contact(run_id, "safe", status="found", email="jobs@acme.ma", level="job_post",
                         url="https://example.test/job", confidence=100, source="linkedin")

    output = export_study(store, run_id, tmp_path / "market.xlsx", include_descriptions=True)

    workbook = load_workbook(output)
    assert workbook.sheetnames == ["IT Jobs Data", "Requirements", "Contacts", "Summary", "Run Info", "Descriptions"]
    jobs = workbook["IT Jobs Data"]
    assert jobs["B2"].value == "'=Safe title"
    assert jobs["L2"].hyperlink.target == "https://example.test/job"
    assert workbook["Descriptions"]["E2"].value == "Description complète"
    assert jobs["M2"].hyperlink.target == "#'Descriptions'!E2"
    assert jobs.freeze_panes == "C2"
    assert jobs.auto_filter.ref is None
    assert jobs.tables['Jobs'].autoFilter.ref == jobs.dimensions
    assert workbook["Requirements"]["E2"].value == "Python"
    assert workbook["Contacts"]["D2"].value == "jobs@acme.ma"
    assert workbook["Contacts"]["G2"].value == 100
    assert jobs["O2"].value == "jobs@acme.ma"
    assert jobs["R2"].value == 100


def test_export_can_limit_output_to_selected_offer_ids(tmp_path):
    store = StudyStore(tmp_path / "studies.db")
    run_id = store.create_run(RunSpec(name="Selected", queries=("python",), cities=("Rabat",), sources=("linkedin",)))
    store.commit_discovery(run_id, "search:1", [
        JobObservation("linkedin", "one", "https://example.test/one", "One"),
        JobObservation("linkedin", "two", "https://example.test/two", "Two"),
    ])
    selected_id = store.offers(run_id)[1]["id"]

    output = export_study(store, run_id, tmp_path / "selected.xlsx", offer_ids=(selected_id,))

    workbook = load_workbook(output)
    assert workbook["IT Jobs Data"].max_row == 2
    assert workbook["IT Jobs Data"]["B2"].value == "Two"
