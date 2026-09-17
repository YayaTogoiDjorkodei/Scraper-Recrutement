"""Load real function bodies without importing the scripts' GUI/network setup.

This is temporary test infrastructure, not a replacement engine. AST compilation
preserves each function body and line number. It avoids refactoring production
initialization merely to create baseline tests. UI tests separately run the full
entry points with a stubbed GeoNames response and a real Tk event loop.
"""

import ast
from contextlib import redirect_stdout
import io
from pathlib import Path
import random
import re
import threading
import types
import unittest
from unittest.mock import Mock, patch
from urllib.parse import quote_plus, urlparse

from bs4 import BeautifulSoup
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd
from rapidfuzz import fuzz
import requests

import liste
from reference_overrides import mots_technicien, niveaux_master
from scraper_config import error_kind

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = Path(__file__).resolve().parent / "fixtures"


def load_functions(source, code=None):
    path = ROOT / source
    module = ast.parse(code if code is not None else path.read_text(encoding="utf-8-sig"), filename=str(path))
    functions = ast.Module(body=[node for node in module.body if isinstance(node, ast.FunctionDef)], type_ignores=[])
    namespace = dict(__name__="legacy_test", random=random, re=re, threading=threading,
        requests=requests, BeautifulSoup=BeautifulSoup, fuzz=fuzz, pd=pd, Font=Font,
        PatternFill=PatternFill, Alignment=Alignment, get_column_letter=get_column_letter,
        quote_plus=quote_plus, urlparse=urlparse, error_kind=error_kind,
        SETTINGS={"proxy_file": "test-proxies.txt", "geonames_username": "test-account"},
        BASE_URL="https://www.linkedin.com/jobs/search/" if source == "main.py" else "https://www.indeed.com/jobs",
        user=types.SimpleNamespace(random="FixtureUserAgent/1.0"), HEADERS={},
        technologies=liste.technologies, niveaux_etudes=liste.niveaux_etudes,
        experience=liste.experience, type_contrat=liste.type_contrat,
        mots_technicien=mots_technicien, niveaux_master=niveaux_master,
        stop_event=threading.Event(), toutes_les_donnees=[], Donner_Exporter=False,
        Statu=Mock(), Bouton_demarrer=Mock(), Bouton_arrete=Mock(), progress=Mock(),
        fenetre=Mock(), messagebox=Mock(), time=Mock(time=Mock(side_effect=[0.0, 1.0]), sleep=Mock()))
    namespace["Bouton_demarrer"].cget.return_value = "normal"
    namespace["fenetre"].after.side_effect = lambda delay, callback: callback()
    exec(compile(functions, str(path), "exec"), namespace)
    return namespace


def fixture(name):
    return (FIXTURES / name).read_text(encoding="utf-8")


def canonical_records(rows):
    # liste.py intentionally contains a set. Ignore only list ordering, not
    # values/duplicates, when comparing snapshots across Python hash seeds.
    return [{key: sorted(value) if isinstance(value, list) else value for key, value in row.items()} for row in rows]


class OfflineTestCase(unittest.TestCase):
    def setUp(self):
        # A forgotten HTTP mock must fail locally, never contact a real source.
        guard = patch("requests.sessions.Session.request", side_effect=AssertionError("Live HTTP is forbidden in tests"))
        guard.start()
        self.addCleanup(guard.stop)
        self.console = io.StringIO()
        output = redirect_stdout(self.console)
        output.__enter__()
        self.addCleanup(output.__exit__, None, None, None)
