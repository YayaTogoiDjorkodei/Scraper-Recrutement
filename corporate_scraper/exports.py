"""Safe Excel exports for a persisted Corporate Scraper study."""

from __future__ import annotations

from collections import Counter
import json
from pathlib import Path
import tempfile
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .store import StudyStore
from .models import FieldEvidence
from .extraction import requirement_summary
from .text import description_text


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LINK_FONT = Font(color="0563C1", underline="single")
FORMULA_PREFIXES = ("=", "+", "-", "@")


def _description_chunks(text: str) -> list[str]:
    # Budget visual lines as well as characters: lists may contain many short
    # lines. Concatenating the chunks must reproduce the entire source text.
    units = [line[start:start + 120] for line in text.splitlines(keepends=True)
             for start in range(0, len(line), 120)]
    return ["".join(units[start:start + 18]) for start in range(0, len(units), 18)]


def _safe(value: object) -> object:
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]", "", value)[:32760]
        if value.lstrip().startswith(FORMULA_PREFIXES):
            return "'" + value
    return value


def _style_table(sheet, name: str, widths: dict[str, int]) -> None:
    sheet.freeze_panes = "C2" if sheet.title == "IT Jobs Data" else "A2"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    sheet.row_dimensions[1].height = 32
    # Table owns its filter. A second sheet AutoFilter over the same range
    # triggers Excel repair even though openpyxl can reopen the file.
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 60 if sheet.title in ("IT Jobs Data", "Requirements") else 30
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if not cell.hyperlink:
                cell.font = Font(name="Calibri", size=11, color="172033")
    if sheet.max_row >= 2:
        table = Table(displayName=name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)
    else:
        sheet.auto_filter.ref = sheet.dimensions


def export_study(store: StudyStore, run_id: str, destination: str | Path, include_descriptions: bool = False,
                 offer_ids: tuple[int, ...] | None = None) -> Path:
    """Atomically write the default workbook and leave an existing file intact on failure."""
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    run = store.run(run_id)
    offers = store.offers(run_id, offer_ids=offer_ids)
    workbook = Workbook()
    jobs = workbook.active
    jobs.title = "IT Jobs Data"
    jobs.append(["Entreprise", "Poste", "Ville / région", "Compétences", "Niveau / diplôme", "Expérience indiquée",
                 "Contrat", "Langues", "Correspondance aux filtres", "Source", "Publication", "Offre originale", "Description", "Identifiant",
                 "Email de contact", "État contact", "Origine email", "Confiance email"])
    requirements = workbook.create_sheet("Requirements")
    requirements.append(["Entreprise", "Poste", "Source", "Champ", "Valeur", "Extrait justificatif", "Méthode", "Score", "Version", "Identifiant"])
    contacts = workbook.create_sheet("Contacts")
    contacts.append(["Identifiant", "Entreprise", "Poste", "Email de contact", "État", "Origine", "Confiance", "Page vérifiée"])
    descriptions = workbook.create_sheet("Descriptions") if include_descriptions else None
    if descriptions is not None:
        descriptions.append(["Identifiant", "Entreprise", "Poste", "Partie", "Description complète"])
    coverage: Counter[str] = Counter()
    qualifications: Counter[str] = Counter()
    company_counts: Counter[str] = Counter()
    city_counts: Counter[str] = Counter()
    source_counts: Counter[str] = Counter()
    skill_counts: Counter[str] = Counter()
    for offer in offers:
        saved_evidence = store.evidence_for_offer(offer["id"])
        evidence = tuple(FieldEvidence(item["field"], item["value"], item["excerpt"], item["method"], item["score"], item["rule_version"])
                         for item in saved_evidence)
        fields = requirement_summary(evidence)
        skills = fields["skills"] or tuple(json.loads(offer["skills_json"]))
        experience = fields["experience"] or offer["experience"]
        education = fields["education"] or offer["education"]
        contract = fields["contract"] or offer["contract"]
        text = description_text(offer["description"] or "")
        identifier = f"{offer['source']}:{offer['source_key']}"
        qualification = {"matching": "Correspond", "excluded": "Hors critères", "unknown": "À vérifier"}.get(offer["qualification_status"], offer["qualification_status"])
        row = [_safe(offer["company"] or "Unknown"), _safe(offer["title"]), _safe(offer["location"] or "Unknown"),
               _safe(offer["source"]), "Ouvrir l'offre", _safe(offer["posted_text"] or "Non indiqué"),
               _safe(", ".join(skills) or "Non indiqué"), _safe(experience or "Non indiquée"), _safe(contract or "Non indiqué"),
               "Récupérée" if text else "Indisponible", qualification, _safe(education or "Non indiqué"),
               _safe(fields["languages"] or "Non indiquées"), identifier]
        jobs.append([row[index] for index in (0, 1, 2, 6, 11, 7, 8, 12, 10, 3, 5, 4, 9, 13)])
        jobs.cell(jobs.max_row, 15, _safe(offer["contact_email"] or "Non trouvé"))
        jobs.cell(jobs.max_row, 16, _safe(offer["contact_status"] or "not_requested"))
        jobs.cell(jobs.max_row, 17, _safe(offer["contact_level"] or "Non vérifié"))
        jobs.cell(jobs.max_row, 18, offer["contact_confidence"] if offer["contact_confidence"] is not None else "Non évaluée")
        contacts.append([identifier, _safe(offer["company"]), _safe(offer["title"]), _safe(offer["contact_email"] or "Non trouvé"),
                         _safe(offer["contact_status"] or "not_requested"), _safe(offer["contact_level"] or "Non vérifié"),
                         offer["contact_confidence"] if offer["contact_confidence"] is not None else "Non évaluée", _safe(offer["contact_url"] or "")])
        if offer["contact_url"]:
            contacts.cell(contacts.max_row, 8).hyperlink = offer["contact_url"]
            contacts.cell(contacts.max_row, 8).font = LINK_FONT
        if descriptions is not None and text:
            first_row = descriptions.max_row + 1
            chunks = _description_chunks(text)
            for part, chunk in enumerate(chunks, 1):
                descriptions.append([identifier, _safe(offer["company"]), _safe(offer["title"]), part, _safe(chunk)])
            cell = jobs.cell(jobs.max_row, 13)
            cell.value = "Lire le texte"
            cell.hyperlink = f"#'Descriptions'!E{first_row}"
            cell.font = LINK_FONT
        qualifications[qualification] += 1
        for label, present in (("Descriptions récupérées", text), ("Compétences détectées", skills), ("Études détectées", education),
                               ("Expérience détectée", experience), ("Contrat détecté", contract), ("Emails publics trouvés", offer["contact_status"] == "found")):
            coverage[label] += bool(present)
        link_cell = jobs.cell(jobs.max_row, 12)
        if offer["canonical_url"]:
            link_cell.hyperlink = offer["canonical_url"]
            link_cell.font = LINK_FONT
        company_counts[offer["company"] or "Unknown"] += 1
        city_counts[offer["location"] or "Unknown"] += 1
        source_counts[offer["source"]] += 1
        skill_counts.update(skills)
        for item in saved_evidence:
            requirements.append([_safe(offer["company"] or "Unknown"), _safe(offer["title"]), _safe(offer["source"]),
                                 _safe(item["field"]), _safe(item["value"]), _safe(description_text(item["excerpt"])), _safe(item["method"]),
                                 item["score"], _safe(item["rule_version"]), identifier])
    _style_table(jobs, "Jobs", {"A": 26, "B": 38, "C": 23, "D": 42, "E": 28, "F": 34, "G": 20, "H": 23, "I": 25, "J": 14, "K": 20, "L": 19, "M": 18, "N": 32, "O": 35, "P": 18, "Q": 22, "R": 18})
    _style_table(requirements, "RequirementsTable", {"A": 26, "B": 36, "C": 14, "D": 23, "E": 32, "F": 105, "G": 14, "H": 12, "I": 14, "J": 32})
    _style_table(contacts, "ContactsTable", {"A": 32, "B": 26, "C": 36, "D": 36, "E": 18, "F": 28, "G": 14, "H": 60})
    if descriptions is not None:
        _style_table(descriptions, "DescriptionTable", {"A": 32, "B": 26, "C": 36, "D": 10, "E": 125})
        for index in range(2, descriptions.max_row + 1):
            descriptions.row_dimensions[index].height = 300

    summary = workbook.create_sheet("Summary")
    summary.append(["Groupe", "Valeur", "Nombre d'offres"])
    summary.append(["Étude", "Offres exportées", len(offers)])
    for group, counts in (("Couverture", coverage), ("Qualification", qualifications), ("Entreprise", company_counts), ("Ville", city_counts), ("Source", source_counts), ("Compétence", skill_counts)):
        for value, count in counts.most_common():
            summary.append([group, _safe(value), count])
    _style_table(summary, "SummaryTable", {"A": 18, "B": 36, "C": 12})

    info = workbook.create_sheet("Run Info")
    info.append(["Field", "Value"])
    info.append(["Study ID", run_id])
    info.append(["Study name", _safe(run["name"])])
    info.append(["State", _safe(run["state"])])
    info.append(["Created at", run["created_at"]])
    info.append(["Updated at", run["updated_at"]])
    info.append(["Run specification", _safe(run["spec_json"])])
    info.append(["Lecture", "Non indiqué = aucune valeur explicite détectée. Vérifier l'extrait dans Requirements. Un intitulé Ingénieur ne prouve pas un diplôme."])
    info.append(["Périmètre", "Les offres Hors critères et À vérifier sont conservées. La colonne de correspondance expose le résultat des filtres."])
    info.append(["Descriptions", "Texte complet dans Descriptions, réparti en parties pour conserver les descriptions longues. Identifiant commun à toutes les feuilles."])
    info.append(["Contacts", "Optionnelle : offre, profil LinkedIn public du recruteur, profil LinkedIn de l’entreprise, site officiel puis page Contact explicitement liée. Le score indique la confiance liée à la source publique, pas la confirmation qu’une personne répondra. Aucune adresse n’est devinée."])
    _style_table(info, "RunInfoTable", {"A": 24, "B": 100})
    for index in range(2, info.max_row + 1):
        info.row_dimensions[index].height = 60
    # Keep the overview before the larger supporting text worksheet.
    if descriptions is not None:
        workbook.move_sheet(descriptions, offset=2)

    with tempfile.NamedTemporaryFile(delete=False, dir=destination.parent, prefix=f".{destination.stem}.", suffix=".xlsx") as handle:
        temporary = Path(handle.name)
    try:
        workbook.save(temporary)
        temporary.replace(destination)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise
    return destination
